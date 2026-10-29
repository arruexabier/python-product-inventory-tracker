"""Command-line interface for the product inventory tracker."""

from __future__ import annotations

import argparse
import csv
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from inventory_tracker.config import load_config

from inventory_tracker.compare import CHANGE_FIELDS, compare_snapshots
from inventory_tracker.scraper import scrape_catalogue

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRODUCT_FIELDS = [
    "product_url",
    "title",
    "upc",
    "sku",
    "brand",
    "price_gbp",
    "price_excl_tax",
    "price_incl_tax",
    "tax",
    "availability",
    "rating",
    "category",
    "product_type",
    "review_count",
    "description",
    "scraped_at",
]
FAILURE_FIELDS = [
    "product_url",
    "title",
    "stage",
    "error",
]


def read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", newline="", encoding="utf-8-sig") as file:
        return list(csv.DictReader(file))


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

def write_xlsx(
    path: Path,
    rows: list[dict[str, Any]],
    fieldnames: list[str],
    sheet_name: str = "Inventory",
) -> None:
    """Write records to a formatted Excel workbook."""
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    worksheet.sheet_view.showGridLines = False

    # Write the column headings.
    worksheet.append(fieldnames)

    # Write the data rows.
    for row_number, row in enumerate(rows, start=2):
        for column_number, field in enumerate(fieldnames, start=1):
            value = row.get(field, "")

            if value is None:
                value = ""

            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            # Keep scraped text beginning with "=" as text,
            # rather than allowing Excel to interpret it as a formula.
            if isinstance(value, str) and value.startswith("="):
                cell.value = value
                cell.data_type = "s"
            else:
                cell.value = value

            # Make product URLs clickable in Excel.
            if (
                field == "product_url"
                and isinstance(value, str)
                and value.startswith(("https://", "http://"))
            ):
                cell.hyperlink = value
                cell.style = "Hyperlink"

            # Format numeric price fields as GBP.
            price_fields = {
                "price_gbp",
                "price_excl_tax",
                "price_incl_tax",
                "tax",
                "old_price_gbp",
                "new_price_gbp",
            }

            if (
                field in price_fields
                and isinstance(value, (int, float))
                and not isinstance(value, bool)
            ):
                cell.number_format = '£#,##0.00'

            # Wrap long descriptions and align data at the top.
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(field == "description"),
            )

    # Style the heading row.
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    worksheet.row_dimensions[1].height = 28
    worksheet.freeze_panes = "A2"

    # Add Excel's built-in column filters.
    worksheet.auto_filter.ref = worksheet.dimensions

    # Set readable column widths, using a bounded sample of the data.
    fixed_widths = {
        "product_url": 48,
        "title": 32,
        "description": 45,
    }

    for column_number, field in enumerate(fieldnames, start=1):
        if field in fixed_widths:
            width = fixed_widths[field]
        else:
            lengths = [len(field)]

            for row in rows[:100]:
                value = row.get(field, "")
                lengths.append(len(str(value or "")))

            width = min(max(max(lengths) + 2, 12), 30)

        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = width

    workbook.save(path)



def run(args: argparse.Namespace) -> int:

    try:
        config = load_config(args.config)
    except ValueError as exc:
        print(
            f"ERROR: Invalid scraper configuration: {exc}",
            file=sys.stderr,
        )
        return 1

    output_dir = Path(args.output_dir).expanduser().resolve()
    snapshots_dir = output_dir / "snapshots"
    reports_dir = output_dir / "reports"

    snapshots_dir.mkdir(parents=True, exist_ok=True)
    reports_dir.mkdir(parents=True, exist_ok=True)

    print(f"Starting catalogue scrape: {config.site_name}")

    if args.max_pages is not None:
        print(f"Page limit: {args.max_pages}")
    else:
        print("Page limit: none (follow pagination to the end)")

    # Scrape the catalogue and collect any detail-page failures.
    try:
        (
            products,
            pages_visited,
            duplicates_skipped,
            detail_failures,
        ) = scrape_catalogue(
            max_pages=args.max_pages,
            delay=args.delay,
            timeout=args.timeout,
            max_retries=args.retries,
            config=config,
        )
    except (RuntimeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        print(
            "No new inventory snapshot was saved.",
            file=sys.stderr,
        )
        return 1

    if not products:
        print(
            "ERROR: no products were collected; "
            "refusing to overwrite the inventory.",
            file=sys.stderr,
        )
        return 1

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")

    # Save a CSV report of any failed product-detail requests.
    if detail_failures:
        failure_path = reports_dir / f"detail_failures_{stamp}.csv"
        write_csv(failure_path, detail_failures, FAILURE_FIELDS)

    # Limited runs are for testing, not inventory comparisons.
    # They must not overwrite the current inventory or create snapshots.
    if args.max_pages is not None:
        preview_csv = output_dir / "preview_inventory.csv"
        preview_xlsx = output_dir / "preview_inventory.xlsx"

        write_csv(preview_csv, products, PRODUCT_FIELDS)
        write_xlsx(
            preview_xlsx,
            products,
            PRODUCT_FIELDS,
            sheet_name="Preview",
        )

        print("\nPreview run completed")
        print(f"  Pages visited:      {pages_visited}")
        print(f"  Products collected: {len(products)}")
        print(f"  Duplicates skipped: {duplicates_skipped}")
        print(f"  Detail failures:    {len(detail_failures)}")
        print(f"  Preview CSV:        {preview_csv}")
        print(f"  Preview Excel:      {preview_xlsx}")
        print("  Historical comparison: skipped")
        print("  Full inventory snapshot: not created")

        if detail_failures:
            print(
                "  Failure report:     "
                f"{reports_dir / f'detail_failures_{stamp}.csv'}"
            )

        return 0

    # Only a complete catalogue run reaches this point.
    # Use the latest previous complete snapshot as the baseline.
    previous_files = sorted(snapshots_dir.glob("inventory_*.csv"))
    previous_path = previous_files[-1] if previous_files else None
    previous = read_csv(previous_path) if previous_path else []

    current_csv = output_dir / "current_inventory.csv"
    current_xlsx = output_dir / "current_inventory.xlsx"

    snapshot_csv = snapshots_dir / f"inventory_{stamp}.csv"

    changes_csv = reports_dir / f"changes_{stamp}.csv"
    changes_xlsx = reports_dir / f"changes_{stamp}.xlsx"

    # Compare only after the complete catalogue has been collected.
    changes = (
        compare_snapshots(previous, products)
        if previous_path
        else []
    )

    # Keep CSV as the format used for historical snapshots.
    write_csv(current_csv, products, PRODUCT_FIELDS)
    write_csv(snapshot_csv, products, PRODUCT_FIELDS)
    write_csv(changes_csv, changes, CHANGE_FIELDS)

    # Also create convenient Excel files for the user.
    write_xlsx(
        current_xlsx,
        products,
        PRODUCT_FIELDS,
        sheet_name="Inventory",
    )
    write_xlsx(
        changes_xlsx,
        changes,
        CHANGE_FIELDS,
        sheet_name="Changes",
    )

    print("\nFull catalogue run completed")
    print(f"  Pages visited:       {pages_visited}")
    print(f"  Unique products:     {len(products)}")
    print(f"  Duplicate products:  {duplicates_skipped}")
    print(f"  Detail-page failures:{len(detail_failures):>5}")
    print(f"  Inventory CSV:       {current_csv}")
    print(f"  Inventory Excel:     {current_xlsx}")
    print(f"  Saved snapshot:      {snapshot_csv}")
    print(f"  Change report CSV:   {changes_csv}")
    print(f"  Change report Excel: {changes_xlsx}")

    if detail_failures:
        print("\nSome product detail pages could not be processed.")

        for failure in detail_failures[:5]:
            print(
                f"  - {failure['title']}: {failure['error']}"
            )

        if len(detail_failures) > 5:
            print(
                f"  ... and {len(detail_failures) - 5} more."
            )

        print(
            "  See the detail_failures CSV in the reports folder."
        )

    if previous_path:
        counts: dict[str, int] = {}

        for change in changes:
            kind = str(change["change_type"])
            counts[kind] = counts.get(kind, 0) + 1

        if counts:
            print("\nChanges since previous snapshot:")

            for kind, count in sorted(counts.items()):
                print(f"  {kind.replace('_', ' ')}: {count}")
        else:
            print("\nNo changes detected since the previous snapshot.")
    else:
        print("\nNo previous baseline; initial snapshot created.")

    return 0





def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Scrape a demo product catalogue and track inventory changes."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)
    run_parser = subparsers.add_parser("run", help="Scrape the catalogue and compare with the previous run")
    run_parser.add_argument(
        "--max-pages", type=int, default=None,
        help="limit the number of catalogue pages (use 2 for a quick test; default: all pages)",
    )
    run_parser.add_argument(
        "--delay", type=float, default=0.5,
        help="seconds to wait between page requests (minimum 0.2; default: 0.5)",
    )
    run_parser.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    run_parser.add_argument("--retries", type=int, default=3, help="maximum attempts per page")
    run_parser.add_argument("--output-dir", default="data", help="directory for inventory, snapshots and reports")
    run_parser.set_defaults(func=run)

    default_config_path = (
        Path(__file__).resolve().parent.parent / "config.json"
    )

    run_parser.add_argument(
        "--config",
        default=str(default_config_path),
        help="path to the website configuration JSON file",
    )


    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    if args.command == "run":
        if args.max_pages is not None and args.max_pages < 1:
            parser.error("--max-pages must be at least 1")
        if args.delay < 0.2:
            parser.error("--delay must be at least 0.2 seconds")
        if args.timeout <= 0:
            parser.error("--timeout must be positive")
        if args.retries < 1:
            parser.error("--retries must be at least 1")
    return args.func(args)
